# -*- coding: utf-8 -*-
"""
be_regneark.py - energirammeberegning med Social- og Boligstyrelsens regneark for den nye
beregningsmetode (bilag til bygningsreglementet, 2026).

Regnearket er metodens referenceberegning. Vi regner ikke selv energirammen: vi fylder
regnearkets indtastningsfelter ud fra modellen og lader regnearket regne (Excel eller
LibreOffice). Så er det styrelsens formler, der giver resultatet.

    from be_regneark import hent_skabelon, udfyld, genberegn, resultat
    skabelon = hent_skabelon("mappe")                     # henter regnearket fra sbst.dk
    udfyld(skabelon, "projekt.xlsx", data)                # data: se udfyld()
    genberegn("projekt.xlsx")                             # Excel (Windows) eller LibreOffice
    resultat("projekt.xlsx")                              # dict med energiramme og -behov

Regnearket ejes af BUILD/SBi og ligger ikke i git; hent_skabelon() henter det.
"""
from __future__ import division, unicode_literals

import os
import shutil
import subprocess
import tempfile

SKABELON_URL = ("https://www.sbst.dk/Media/639063199880292213/"
                "Be05_ver2026_02_02_Speciel_bygning_til_Metodebeskrivelse.xlsx")
SKABELON_NAVN = "Be05_ver2026_02_02_Speciel_bygning_til_Metodebeskrivelse.xlsx"

# Rækker med indtastning i regnearket (version 02.02.2026)
KONST_RAEKKER = range(6, 26)         # 20 flader
LINJE_RAEKKER = range(30, 40)        # 10 linjetab
VINDUE_RAEKKER = range(6, 26)        # 20 vinduer/døre
SKYGGE_RAEKKER = range(7, 17)        # profil 1..10 (række 6 = default)
VENT_RAEKKER = range(7, 27)          # 20 ventilationszoner (vinter); sommer i række +27
INTERN_RAEKKER = range(6, 26)        # 20 benyttelseszoner; belysning i række +26

ORIENTERINGER = ["n", "nø", "ø", "sø", "s", "sv", "v", "nv"]

# Forsyning: hovedvalg i fanebladet Hoved. Detaljer (effekt, COP osv.) står i regnearket og
# kan ændres med 'celler'.
FORSYNING = {
    "fjernvarme": {"Hoved!F34": "F", "Hoved!F37": "N", "Hoved!F38": "N", "Hoved!F39": "N",
                   "Hoved!F40": "N", "Hoved!F41": "N", "Hoved!F42": "N", "Hoved!F43": "N",
                   "Hoved!F44": "N", "Hoved!F49": 0},
    "luft_vand_varmepumpe": {"Hoved!F34": "E", "Hoved!F37": "N", "Hoved!F38": "N", "Hoved!F39": "N",
                             "Hoved!F40": "K", "Hoved!F41": "N", "Hoved!F42": "N", "Hoved!F43": "N",
                             "Hoved!F44": "N", "Hoved!F49": 0},
    "regneark": {},      # behold regnearkets eksempel uændret
}


def orientering(azimut):
    """Kompasretning i grader (0 = nord, 90 = øst) -> regnearkets kode (n, nø, ø, ...)."""
    return ORIENTERINGER[int(((azimut % 360) + 22.5) // 45) % 8]


def haeldning(grader):
    """Hældning afrundet til regnearkets trin (0, 15, ..., 90)."""
    return int(min(90, max(0, 15 * round(grader / 15.0))))


def hent_skabelon(mappe):
    """Henter regnearket fra Social- og Boligstyrelsen til mappe (kun første gang)."""
    sti = os.path.join(mappe, SKABELON_NAVN)
    if not os.path.isfile(sti):
        if not os.path.isdir(mappe):
            os.makedirs(mappe)
        from urllib.request import urlopen
        data = urlopen(SKABELON_URL, timeout=60).read()
        with open(sti, "wb") as f:
            f.write(data)
    return sti


def _ryd(ws, raekker, kolonner):
    for r in raekker:
        for k in kolonner:
            ws["%s%d" % (k, r)].value = None


def _skriv_raekker(ws, raekker, poster, kolonner, navn):
    raekker = list(raekker)
    if len(poster) > len(raekker):
        raise ValueError("%s: %d rækker, regnearket har plads til %d. Saml flere i samme række."
                         % (navn, len(poster), len(raekker)))
    for r, post in zip(raekker, poster):
        for k, felt in kolonner.items():
            v = post.get(felt)
            if v is not None:
                ws["%s%d" % (k, r)].value = v


def udfyld(skabelon, ud, data):
    """Skriver data ind i en kopi af regnearket.

    data = {
      "bygning": {"navn", "type" (F/S/E/A), "etageareal", "bebygget", "varmekapacitet",
                  "brugstid", "boligenheder"},
      "konstruktioner": [{"navn", "areal", "u", "b", "inde", "ude"}],     # flader
      "linjetab": [{"navn", "laengde", "psi", "b", "inde", "ude"}],
      "vinduer": [{"navn", "antal", "orientering", "haeldning", "areal", "u", "b", "ff", "g",
                   "skygge", "fc"}],
      "skygger": [{"navn", "horisont", "udhaeng", "venstre", "hoejre", "vindueshul"}],  # profil 1..
      "ventilation": [{"navn", "areal", "fo", "qvm", "hvgv", "tind", "elvf", "qid", "qis", "sel",
                       "sommer_qvm", "sommer_qid", "nat_qvm", "nat_qis", "koeling"}],
      "intern": [{"navn", "areal", "personer", "udstyr", "udstyr_nat"}],
      "forsyning": "fjernvarme" | "luft_vand_varmepumpe" | "regneark",
      "celler": {"Hoved!C21": 80, ...}       # vilkårlige felter, skrives til sidst
    }
    """
    import openpyxl
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(skabelon)

    b = data.get("bygning") or {}
    h = wb["Hoved"]
    for celle, felt in (("C7", "navn"), ("C10", "type"), ("C15", "boligenheder"), ("C18", "etageareal"),
                        ("C20", "bebygget"), ("C21", "varmekapacitet"), ("C22", "brugstid")):
        if b.get(felt) is not None:
            h[celle].value = b[felt]

    if "konstruktioner" in data:
        ws = wb["Konst"]
        _ryd(ws, KONST_RAEKKER, "BDEFIJ")
        _skriv_raekker(ws, KONST_RAEKKER, data["konstruktioner"],
                       {"B": "navn", "D": "areal", "E": "u", "F": "b", "I": "inde", "J": "ude"}, "Konstruktioner")
    if "linjetab" in data:
        ws = wb["Konst"]
        _ryd(ws, LINJE_RAEKKER, "BDEFIJ")
        _skriv_raekker(ws, LINJE_RAEKKER, data["linjetab"],
                       {"B": "navn", "D": "laengde", "E": "psi", "F": "b", "I": "inde", "J": "ude"}, "Linjetab")
    if "vinduer" in data:
        ws = wb["Vinduer"]
        _ryd(ws, VINDUE_RAEKKER, "BCDEFGHKLMN")
        _skriv_raekker(ws, VINDUE_RAEKKER, data["vinduer"],
                       {"B": "navn", "C": "antal", "D": "orientering", "E": "haeldning", "F": "areal",
                        "G": "u", "H": "b", "K": "ff", "L": "g", "M": "skygge", "N": "fc"}, "Vinduer")
    # tomme vinduesrækker skal pege på skyggeprofil 0, ellers giver LibreOffice #N/A (Excel tåler det)
    for r in VINDUE_RAEKKER:
        if wb["Vinduer"]["M%d" % r].value is None:
            wb["Vinduer"]["M%d" % r].value = 0
    if "skygger" in data:
        ws = wb["Skygger"]
        _ryd(ws, SKYGGE_RAEKKER, "BCDEFG")
        _skriv_raekker(ws, SKYGGE_RAEKKER, data["skygger"],
                       {"B": "navn", "C": "horisont", "D": "udhaeng", "E": "venstre", "F": "hoejre",
                        "G": "vindueshul"}, "Skygger")
    if "ventilation" in data:
        ws = wb["Vent"]
        _ryd(ws, VENT_RAEKKER, "BCDEFGHIJK")
        _ryd(ws, [r + 27 for r in VENT_RAEKKER], "DEJKO")
        _skriv_raekker(ws, VENT_RAEKKER, data["ventilation"],
                       {"B": "navn", "C": "areal", "D": "fo", "E": "qvm", "F": "hvgv", "G": "tind", "H": "elvf",
                        "I": "qid", "J": "qis", "K": "sel"}, "Ventilation")
        _skriv_raekker(ws, [r + 27 for r in VENT_RAEKKER], data["ventilation"],
                       {"D": "sommer_qvm", "E": "sommer_qid", "J": "nat_qvm", "K": "nat_qis", "O": "koeling"},
                       "Ventilation (sommer)")
    if "intern" in data:
        ws = wb["Intern"]
        _ryd(ws, INTERN_RAEKKER, "BDEFG")
        _skriv_raekker(ws, INTERN_RAEKKER, data["intern"],
                       {"B": "navn", "D": "areal", "E": "personer", "F": "udstyr", "G": "udstyr_nat"}, "Intern")

    celler = dict(FORSYNING[data.get("forsyning") or "regneark"])
    celler.update(data.get("celler") or {})
    for adr, v in celler.items():
        ark, celle = adr.split("!")
        wb[ark][celle].value = v
    wb.save(ud)
    return ud


# --- genberegning --------------------------------------------------------------
def _soffice():
    for k in ("soffice", "libreoffice"):
        p = shutil.which(k)
        if p:
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.isfile(p):
            return p
    return None


def _genberegn_excel(fil):
    import win32com.client                  # pywin32 (Windows med Excel)
    xl = win32com.client.DispatchEx("Excel.Application")
    xl.Visible = False
    xl.DisplayAlerts = False
    try:
        wb = xl.Workbooks.Open(os.path.abspath(fil))
        xl.CalculateFull()
        wb.Save()
        wb.Close(False)
    finally:
        xl.Quit()
    return fil


def _genberegn_libreoffice(fil, soffice):
    """LibreOffice uden brugerflade, med 'genberegn altid ved indlæsning'."""
    profil = tempfile.mkdtemp(prefix="be_lo_")
    os.makedirs(os.path.join(profil, "user"))
    with open(os.path.join(profil, "user", "registrymodifications.xcu"), "w") as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n'
                '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                'xmlns:xs="http://www.w3.org/2001/XMLSchema" '
                'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
                '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
                '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n'
                '</oor:items>\n')
    udmappe = tempfile.mkdtemp(prefix="be_ud_")
    try:
        url = "file:///" + profil.replace("\\", "/").lstrip("/")
        subprocess.run([soffice, "-env:UserInstallation=" + url, "--headless", "--convert-to",
                        "xlsx:Calc MS Excel 2007 XML", "--outdir", udmappe, os.path.abspath(fil)],
                       check=True, timeout=900, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        shutil.copyfile(os.path.join(udmappe, os.path.basename(fil)), fil)
    finally:
        shutil.rmtree(profil, ignore_errors=True)
        shutil.rmtree(udmappe, ignore_errors=True)
    return fil


def genberegn(fil):
    """Genberegner regnearket: Excel (Windows), ellers LibreOffice. Returnerer metoden."""
    try:
        _genberegn_excel(fil)
        return "Excel"
    except Exception:
        pass
    so = _soffice()
    if so:
        _genberegn_libreoffice(fil, so)
        return "LibreOffice"
    raise RuntimeError("Kan ikke genberegne: hverken Excel (pywin32) eller LibreOffice fundet. "
                       "Åbn %s i Excel, gem den, og læs resultatet igen." % fil)


# --- resultat -------------------------------------------------------------------
def resultat(fil):
    """Læser fanebladet RESULTAT (efter genberegning)."""
    import openpyxl
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(fil, data_only=True)["RESULTAT"]
    v = lambda c: ws[c].value
    if not isinstance(v("H30"), (int, float)):
        raise RuntimeError("Regnearket er ikke genberegnet (RESULTAT!H30 = %r)." % (v("H30"),))
    rammer = {}
    for start, navn in ((24, "BR18"), (33, "Lavenergi"), (6, "Renoveringsklasse 2"), (15, "Renoveringsklasse 1")):
        rammer[navn] = {"ramme_uden_tillaeg": v("H%d" % (start + 2)), "tillaeg": v("H%d" % (start + 3)),
                        "ramme": v("H%d" % (start + 4)), "behov": v("H%d" % (start + 6)),
                        "opfyldt": (v("H%d" % start) or "").strip() == "+"}
    bidrag, gruppe = {}, ""
    for r in range(42, 72):
        if v("A%d" % r) and not v("B%d" % r) and not isinstance(v("G%d" % r), (int, float)):
            gruppe = v("A%d" % r).strip()            # fx "Netto behov"
            continue
        etiket = v("B%d" % r) or v("A%d" % r)
        if etiket and isinstance(v("G%d" % r), (int, float)):
            etiket = etiket.strip()
            bidrag["%s: %s" % (gruppe, etiket) if v("B%d" % r) and gruppe else etiket] = v("G%d" % r)
    maaneder = ["jan", "feb", "mar", "apr", "maj", "jun", "jul", "aug", "sep", "okt", "nov", "dec"]
    kol = "EFGHIJKLMNOP"
    pr_maaned = [v("%s106" % k) for k in kol]
    return {"etageareal_m2": v("O8"), "energirammer": rammer, "noegletal_kWh_m2": bidrag,
            "energibehov_pr_maaned_kWh_m2": dict(zip(maaneder, pr_maaned))}
