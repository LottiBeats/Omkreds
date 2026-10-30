"""
Tekstblokkens formatering (**fed**, *kursiv*) — rich_text.py.

Det vigtigste her er, hvad der IKKE bliver til kursiv: regnestykker med
gangetegn og tekst uden markering skal komme uændret igennem til PDF'en.
"""
import rich_text as rt


def test_plain_text_is_unchanged():
    s = "Almindelig tekst\n• punkt\n  1. Lastgrundlag"
    assert rt.to_reportlab(s) == s


def test_bold_and_italic():
    assert rt.to_reportlab("Det er **vigtigt** og *kursivt*.") == \
        "Det er <b>vigtigt</b> og <i>kursivt</i>."
    assert rt.to_reportlab("***begge***") == "<b><i>begge</i></b>"


def test_multiplication_is_not_italic():
    for s in ["1,35*G_k*1,5", "a * b * c", "2*3*4 = 24"]:
        assert rt.to_reportlab(s) == s


def test_docx_runs():
    from docx import Document
    doc = Document()
    p = doc.add_paragraph()
    rt.add_docx_runs(p, "Bemærk: **vigtigt** og *kursivt*")
    runs = [(r.text, bool(r.bold), bool(r.italic)) for r in p.runs]
    assert runs == [("Bemærk: ", False, False), ("vigtigt", True, False),
                    (" og ", False, False), ("kursivt", False, True)]


def test_understregning():
    from rich_text import to_reportlab, runs
    assert to_reportlab("a ++b++ c") == "a <u>b</u> c"
    assert to_reportlab("++**fed**++ og *k*") == "<u><b>fed</b></u> og <i>k</i>"
    assert to_reportlab("C++ og x++ er ikke understreget") == "C++ og x++ er ikke understreget"
    assert runs("++a++ b") == [("a", False, False, True), (" b", False, False, False)]


def test_understregning_i_word():
    from docx import Document
    from rich_text import add_docx_runs
    p = Document().add_paragraph()
    add_docx_runs(p, "x ++u++ **f**")
    assert [(r.text, bool(r.underline), bool(r.bold)) for r in p.runs] == [("x ", False, False), ("u", True, False), (" ", False, False), ("f", False, True)]
