"""Projektets regnearksbog i PDF og Word: ark_visning-blokke peger på et ark
i project["bog"] og evt. et område af det."""
import io

import pytest

from pdf_builder import udfold_arkvisninger, _laes_omraade

BOG = {"ark": [{
    "id": "ark_1", "navn": "Fundament", "titel": "Punktfundament", "overskrift": True,
    "celler": [["Størrelse", "Værdi"], ["V_d", "=B3*2"], ["V_G", "29,96"], ["η", "0,200"]],
    "_vist": [["Størrelse", "Værdi", "Enhed"], ["V_d", "59,92", "kN"], ["V_G", "29,96", "kN"], ["η", "0,200", ""]],
    "bredder": [200, 80, 60],
}]}


def test_laes_omraade():
    assert _laes_omraade("A1:C4") == (0, 0, 3, 2)
    assert _laes_omraade("$B$2:A1") == (0, 0, 1, 1)
    assert _laes_omraade("B3") == (2, 1, 2, 1)
    assert _laes_omraade("") is None and _laes_omraade("x") is None


def test_hele_arket_og_titel_fra_arket():
    [b] = udfold_arkvisninger([{"id": 1, "type": "ark_visning", "data": {"ark": "ark_1", "omraade": "", "titel": ""}}], BOG)
    assert b["type"] == "regneark" and b["data"]["titel"] == "Punktfundament"
    assert b["data"]["_vist"][1] == ["V_d", "59,92", "kN"] and b["data"]["overskrift"] is True


def test_omraade_beskaerer_og_overskrift_kun_fra_foerste_raekke():
    [b] = udfold_arkvisninger([{"id": 1, "type": "ark_visning", "data": {"ark": "ark_1", "omraade": "A2:B3", "titel": "Laster"}}], BOG)
    assert b["data"]["_vist"] == [["V_d", "59,92"], ["V_G", "29,96"]]
    assert b["data"]["bredder"] == [200, 80] and b["data"]["overskrift"] is False
    assert b["data"]["titel"] == "Laster"


def test_slettet_ark_udelades_og_andre_blokke_roeres_ikke():
    tekst = {"id": 2, "type": "text", "data": {"text": "x"}}
    ud = udfold_arkvisninger([{"id": 1, "type": "ark_visning", "data": {"ark": "findes_ikke"}}, tekst], BOG)
    assert ud == [tekst]
    assert udfold_arkvisninger([tekst], None) == [tekst]


def test_pdf_og_word_viser_arket():
    pdfium = pytest.importorskip("pypdfium2")
    from pdf_builder import build_pdf, build_pdf_simpel
    blokke = [{"id": 1, "type": "ark_visning", "data": {"ark": "ark_1", "omraade": "A1:C2"}}]
    projekt = {"id": "x", "metadata": {"project_name": "Test"}, "documents": {}, "bog": BOG}
    for pdf in (build_pdf(projekt, blokke, doc_id="A2"), build_pdf_simpel(blokke, bog=BOG)):
        t = "\n".join(p.get_textpage().get_text_range() for p in pdfium.PdfDocument(pdf))
        assert "59,92" in t and "Punktfundament" in t
        assert "29,96" not in t, "kun området A1:C2 vises"
    pytest.importorskip("docx")
    import docx as d
    from word_builder import build_word
    w = build_word(projekt, blokke, doc_id="A2")
    t = "\n".join(c.text for tb in d.Document(io.BytesIO(w)).tables for r in tb.rows for c in r.cells)
    assert "59,92" in t and "29,96" not in t


def test_raekketyper_som_i_frontend():
    from pdf_builder import _regneark_raekketyper, _celle_art
    rows = [["Lag", "G"], ["TAG — pr. m²", ""], ["Tegl", "0,5"], ["Σ tag", "0,5"], ["", ""], ["RESULTATER", "Værdi"], ["Noter", ""], ["NOTE", ""]]
    fed = {"A2": True, "A6": True, "B6": True}
    art = [[_celle_art(c) for c in r] for r in rows]
    assert _regneark_raekketyper(rows, fed, art, True) == ["hoved", "sektion", "", "sum", "tom", "hoved", "", "sektion"]
    assert [_celle_art(x) for x in ["=A1", "1,5", "Lag", "", "−3"]] == ["f", "i", "t", "", "i"]


def test_pdf_med_formatering_og_overskrift_midt_i_arket():
    pytest.importorskip("pypdfium2")
    from pdf_builder import build_pdf_simpel, _regneark_raekketyper, _celle_art
    rows = [["Rem", ""], ["Bredde B", "4"], ["Last", "q"], ["G sup", "1,38"]]
    art = [[_celle_art(c) for c in r] for r in rows]
    assert _regneark_raekketyper(rows, {"A1": True, "A3": True, "B3": True}, art, False) == ["sektion", "", "hoved", ""]
    bog = {"ark": [{"id": "a", "navn": "A", "overskrift": False, "celler": rows, "_vist": rows,
                    "fmt": {"B4": {"fyld": "#fff2cc", "ramme": {"b": 2}, "kursiv": True}}}]}
    assert len(build_pdf_simpel([{"id": 1, "type": "ark_visning", "data": {"ark": "a"}}], bog=bog)) > 1000
