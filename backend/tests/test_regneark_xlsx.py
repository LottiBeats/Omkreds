"""Regnearksbogen til og fra Excel (regneark_xlsx.py).

Formlerne oversættes mellem dansk Excel (=HVIS(A1>0;1,5;0)) og Excels
interne form (=IF(A1>0,1.5,0)). Den fulde kontrol mod en rigtig regneark-
motor ligger i test_eksport_genberegnes_ens_i_libreoffice (kun hvor
LibreOffice er installeret)."""
import base64
import io
import math
import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from regneark_xlsx import bog_til_xlsx, formel_fra_excel, formel_til_excel, xlsx_til_bog


@pytest.mark.parametrize("da,en", [
    ("=HVIS(A1>0;B2*1,5;0)", "=IF(A1>0,B2*1.5,0)"),
    ("=LOPSLAG(\"Tka\";Opbygninger!$A$2:$L$10;3;FALSK)", "=VLOOKUP(\"Tka\",Opbygninger!$A$2:$L$10,3,FALSE)"),
    ("='Stabilitet V1'!F19*2", "='Stabilitet V1'!F19*2"),
    ("=\"a;b,c\"&TEKST", "=\"a;b,c\"&TEKST"),
    ("=SUMHVIS(A2:A9;\"J1.D503\";L2:L9)", "=SUMIF(A2:A9,\"J1.D503\",L2:L9)"),
    ("=HVIS.FEJL(XOPSLAG(A1;B:B;C:C);0)", "=IFERROR(_xlfn.XLOOKUP(A1,B:B,C:C),0)"),
    ("=KVROD(POTENS(2;3))+RADIANER(45)", "=SQRT(POWER(2,3))+RADIANS(45)"),
    ("=g_kip_V4*10", "=g_kip_V4*10"),
])
def test_formler_begge_veje(da, en):
    assert formel_til_excel(da) == en
    tilbage = formel_fra_excel(en, maks_raekke=40)
    if "B:B" in da:
        assert tilbage == "=HVIS.FEJL(XOPSLAG(A1;B1:B40;C1:C40);0)"
    else:
        assert tilbage == da


def test_ukendte_funktioner_og_sand_falsk():
    ukendte = set()
    assert formel_fra_excel("=IF(TRUE,FALSE(),OFFSET(A1,1,1))", ukendte) == "=HVIS(SAND;FALSK();OFFSET(A1;1;1))"
    assert ukendte == {"OFFSET"}


BOG = {"ark": [
    {"id": "a1", "navn": "Laster", "celler": [["q_p", "1,36", "kN/m²"], ["c", "0,8", ""], ["w", "=B1*B2", "kN/m²"]],
     "bredder": [150, 92, 60], "dec": {"B3": 3}, "fed": {"A3": True}},
    {"id": "a2", "navn": "Stabilitet V1", "celler": [["F", "=Laster!B3*10+w_e"], ["tekst", "SAND"]]},
], "navne": {"w_e": {"ark": "a1", "celle": "B3", "enhed": "kN/m²"}}}


def test_eksport_formater_og_navne():
    wb = load_workbook(io.BytesIO(bog_til_xlsx(BOG)))
    ws = wb["Laster"]
    assert ws["B1"].value == 1.36 and ws["B3"].value == "=B1*B2"
    assert ws["B3"].number_format == "0.000" and ws["A3"].font.b
    assert wb["Stabilitet V1"]["B1"].value == "=Laster!B3*10+w_e"
    assert wb["Stabilitet V1"]["B2"].value is True
    assert list(wb.defined_names["w_e"].destinations) == [("Laster", "$B$3")]


def test_eksport_og_import_igen():
    r = xlsx_til_bog(bog_til_xlsx(BOG))
    assert [a["navn"] for a in r["ark"]] == ["Laster", "Stabilitet V1"]
    assert r["ark"][0]["celler"][0] == ["q_p", "1,36", "kN/m²"]
    assert r["ark"][0]["celler"][2][1] == "=B1*B2"
    assert r["ark"][1]["celler"][0][1] == "=Laster!B3*10+w_e"
    assert r["ark"][1]["celler"][1][1] == "SAND"
    assert r["ark"][0]["dec"]["B3"] == 3 and r["ark"][0]["fed"]["A3"] is True
    assert r["navne"] == {"w_e": {"ark": "Laster", "celle": "B3", "enhed": ""}}
    assert r["advarsler"] == []


def test_ugyldige_filer():
    with pytest.raises(ValueError, match="xlsx"):
        xlsx_til_bog(b"ikke en fil")


def test_endpoints(client):
    r = client.post("/regneark/xlsx", json=BOG)
    assert r.status_code == 200 and r.content[:2] == b"PK"
    r2 = client.post("/regneark/fra-xlsx", json={"data": base64.b64encode(r.content).decode()})
    assert r2.status_code == 200 and r2.json()["ark"][0]["navn"] == "Laster"
    assert client.post("/regneark/fra-xlsx", json={"data": "eA=="}).status_code == 422


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice er ikke installeret")
def test_eksport_genberegnes_ens_i_libreoffice(tmp_path):
    """LibreOffice regner den eksporterede fil; tallene skal være Omkreds' (regnet i hånden her)."""
    bog = {"ark": [
        {"id": "o", "navn": "Opbygninger", "celler": [
            ["ID", "g_inf", "g_sup", "s"], ["TRa", "6", "6,55", "0"], ["Tka", "0,5", "1,5", "0,8"]]},
        {"id": "e", "navn": "Elementer", "celler": [
            ["El", "xA", "xB", "ID", "x1", "x2", "L", "g_sup", "RB"],
            ["J1", "0", "4,2", "TRa", "3", "4,2", "=C2-B2", "=LOPSLAG(D2;Opbygninger!$A$2:$D$3;3;FALSK)",
             "=HVIS(D2=\"\";\"\";H2*(F2-E2)*((E2+F2)/2-B2)/G2)"],
            ["", "", "", "", "", "", "", "=SUMHVIS(A2:A3;\"J1\";I2:I3)", "=RUND.OP(KVROD(2);2)+g_s"]]},
    ], "navne": {"g_s": {"ark": "o", "celle": "C3"}}}
    (tmp_path / "ind.xlsx").write_bytes(bog_til_xlsx(bog))
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp_path / "ud"),
                    str(tmp_path / "ind.xlsx")], check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "ud" / "ind.xlsx", data_only=True)["Elementer"]
    assert math.isclose(ws["H2"].value, 6.55)
    assert math.isclose(ws["I2"].value, 6.55 * 1.2 * 3.6 / 4.2)
    assert math.isclose(ws["H3"].value, 6.55 * 1.2 * 3.6 / 4.2)
    assert math.isclose(ws["I3"].value, 1.42 + 1.5)


def test_formatering_frem_og_tilbage():
    bog = {"ark": [{"id": "a", "navn": "A", "celler": [["x", "1,5"]],
                    "fmt": {"B1": {"fyld": "#fff2cc", "ramme": {"b": 2, "t": 1}, "kursiv": True, "just": "c"}}}]}
    ws = load_workbook(io.BytesIO(bog_til_xlsx(bog)))["A"]
    assert ws["B1"].fill.fgColor.rgb.endswith("FFF2CC") and ws["B1"].border.bottom.style == "medium"
    assert ws["B1"].font.i and ws["B1"].alignment.horizontal == "center"
    assert xlsx_til_bog(bog_til_xlsx(bog))["ark"][0]["fmt"]["B1"] == {"kursiv": True, "fyld": "#fff2cc", "ramme": {"t": 1, "b": 2}, "just": "c"}


def test_skjulte_kolonner_nul_og_formater_paa_tomme_celler():
    """Hjælpekolonner er skjulte, en formel med 0 vises som –, og rammer/fyld står også på tomme celler."""
    import io as _io
    import openpyxl
    bog = {"ark": [{"id": "a", "navn": "Lodret", "celler": [["ID", "", "x"], ["V1", "=0*1", "0"]],
                    "nul": "–", "skjulte": [2], "dec": {"B2": 2},
                    "fmt": {"C1": {"fyld": "#ddebf7"}, "D2": {"ramme": {"t": 2}}, "B2": {"graa": True}}}]}
    ws = openpyxl.load_workbook(_io.BytesIO(bog_til_xlsx(bog)))["Lodret"]
    assert ws.column_dimensions["C"].hidden
    assert ws["B2"].number_format == '0.00;-0.00;"–"'
    assert ws["C2"].number_format == "General"          # indtastet 0 er et tal
    assert ws["C2"].font.color.rgb.endswith("1D4ED8")   # indtastede tal i blåt
    assert ws["B2"].font.color.rgb.endswith("8A8A8A")
    assert ws["D2"].border.top.style == "medium"        # tom celle med ramme
    assert ws.freeze_panes == "A2"
    tilbage = xlsx_til_bog(bog_til_xlsx(bog))["ark"][0]
    assert tilbage["skjulte"] == [2]


def test_udskrift_og_gitter():
    """Arkene udskrives liggende med hele bredden på én side; gitter: false skjuler gitterlinjerne (og læses tilbage)."""
    bog = {"ark": [{"id": "a", "navn": "Lastnedføring", "celler": [["x"]], "gitter": False},
                   {"id": "b", "navn": "Andet", "celler": [["y"]]}]}
    wb = load_workbook(io.BytesIO(bog_til_xlsx(bog)))
    ws = wb["Lastnedføring"]
    assert ws.page_setup.orientation == "landscape" and ws.sheet_properties.pageSetUpPr.fitToPage
    assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
    assert ws.sheet_view.showGridLines is False and wb["Andet"].sheet_view.showGridLines is not False
    tilbage = xlsx_til_bog(bog_til_xlsx(bog))["ark"]
    assert tilbage[0].get("gitter") is False and "gitter" not in tilbage[1]
